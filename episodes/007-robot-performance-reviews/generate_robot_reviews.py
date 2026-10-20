"""Generate individualized Word reviews from fictional or authorized CSV records."""
import argparse
import csv
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import re
import shutil
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor

FIELDS = ('robot_id', 'robot_name', 'department', 'efficiency_rating',
          'accuracy_rating', 'teamwork_rating', 'manager_comments')
RATINGS = FIELDS[3:6]


def load_records(path):
    """Validate the ENTIRE input before creating any output. Invalid rows abort."""
    with path.open(encoding='utf-8-sig', newline='') as stream:
        reader = csv.DictReader(
            stream, strict=True)
        if not reader.fieldnames or set(reader.fieldnames) != set(FIELDS):
            raise ValueError('CSV must contain exactly the documented columns')
        if len(reader.fieldnames) != len(FIELDS):
            raise ValueError('Duplicate CSV columns')
        records, seen = [], set()
        for number, row in enumerate(reader, 2):
            if None in row or any(not row.get(k) or not row[k].strip() for k in FIELDS):
                raise ValueError(f'Row {number}: missing or extra fields')
            row = {k: v.strip() for k, v in row.items()}
            robot_id = row['robot_id']
            if not re.fullmatch(r'R[0-9]{3,6}', robot_id):
                raise ValueError(f'Row {number}: robot_id must be R plus 3 to 6 digits')
            if robot_id in seen:
                raise ValueError(f'Row {number}: duplicate robot_id {robot_id}')
            for key in RATINGS:
                if row[key] not in ('1', '2', '3', '4', '5'):
                    raise ValueError(f'Row {number}: {key} must be an integer from 1 to 5')
                row[key] = int(row[key])
            for key in ('robot_name', 'department', 'manager_comments'):
                if len(row[key]) > 240 or any(ord(c) < 32 for c in row[key]):
                    raise ValueError(f'Row {number}: {key} must be printable and at most 240 characters')
            seen.add(robot_id)
            records.append(row)
    if not records:
        raise ValueError('CSV contains no employee records')
    return records


def average_rating(robot):
    total = sum(robot[key] for key in RATINGS)
    return (Decimal(total) / 3).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)


def create_review(robot):
    document = Document()
    document.add_paragraph(
        'MINIMUM EFFORT LABS')
    document.add_heading(
        'Annual Robot Performance Review', 0)
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.8)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = document.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles['Normal'].font.size = Pt(11)
    name = robot['robot_name']
    robot_id = robot['robot_id']
    department = robot['department']
    document.add_paragraph(
        f"Robot Name: {name}")
    document.add_paragraph(
        f"Robot ID: {robot_id}")
    document.add_paragraph(
        f"Department: {department}")
    score = average_rating(robot)
    document.add_heading('Review Summary', 1)
    document.add_paragraph(f"{robot['robot_name']} earned an overall rating of {score} out of 5.00.")
    table = document.add_table(
        rows=1, cols=2)
    table.style = 'Table Grid'
    table.rows[0].cells[0].text = 'Performance category'
    table.rows[0].cells[1].text = 'Rating out of 5'
    for label, key in zip(('Efficiency', 'Accuracy', 'Teamwork'), RATINGS):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(robot[key])
    cells = table.add_row().cells
    cells[0].text = 'Overall Average'
    cells[1].text = str(score)
    document.add_heading('Manager Comments', 1)
    document.add_paragraph(robot['manager_comments'])
    document.add_heading('Overall Recommendation', 1)
    recommendation = ('Recognize strong performance' if score >= 4 else
                      'Continue development with manager support' if score >= 3 else
                      'Agree a focused improvement plan')
    document.add_paragraph(recommendation)
    document.add_paragraph('Fictional training example only. Not a real employment assessment.')
    return document


def generate_reviews(csv_path, output):
    records = load_records(csv_path)
    # mkdir(exist_ok=False) reserves a NEW directory; no existing files are touched.
    output.mkdir(parents=True, exist_ok=False)
    try:
        for robot in records:
            document = create_review(robot)
            robot_id = robot['robot_id']
            filename = f'{robot_id}_review.docx'
            document.save(output / filename)
    except Exception:
        # Remove only files in the newly reserved directory owned by this run.
        shutil.rmtree(output)
        raise
    return len(records)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('csv_file', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    try:
        count = generate_reviews(args.csv_file, args.output)
    except (OSError, ValueError, csv.Error) as error:
        print(f'Error: {error}', file=sys.stderr)
        return 2
    print(f'Records processed: {count}')
    print(f'Reviews generated: {count}')
    print(f'Output directory: {args.output.as_posix()}/')
    return 0


if __name__ == '__main__':
    sys.exit(main())
