import csv
from decimal import Decimal
import hashlib
import importlib.util
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import generate_robot_reviews as app
from generate_sample_data import generate

class ReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name);self.csv=self.root/'robots.csv';self.out=self.root/'reviews'
        self.row=['R001','Unit Alpha','Accounting','5','4','3','Accurate, careful and reliable.']

    def write(self, rows=None, fields=app.FIELDS):
        with self.csv.open('w',encoding='utf-8',newline='') as stream:
            w=csv.writer(stream);w.writerow(fields);w.writerows(rows if rows is not None else [self.row])

    def invalid(self, rows=None, fields=app.FIELDS):
        self.write(rows,fields)
        with self.assertRaises(ValueError):app.generate_reviews(self.csv,self.out)
        self.assertFalse(self.out.exists())

    def test_valid_personalized_document(self):
        self.write();before=self.csv.read_bytes();self.assertEqual(app.generate_reviews(self.csv,self.out),1)
        d=Document(self.out/'R001_review.docx');text='\n'.join(p.text for p in d.paragraphs)
        for value in ['Unit Alpha','R001','Accounting',self.row[-1],'4.00','Overall Recommendation']:
            self.assertIn(value,text)
        self.assertEqual(d.paragraphs[1].style.name,'Title')
        self.assertEqual(len(d.tables),1);self.assertEqual(d.tables[0].style.name,'Table Grid')
        self.assertEqual([r.cells[1].text for r in d.tables[0].rows],['Rating out of 5','5','4','3','4.00'])
        self.assertEqual(self.csv.read_bytes(),before)

    def test_missing_columns(self):self.invalid([self.row[:-1]],app.FIELDS[:-1])
    def test_duplicate_columns(self):self.invalid([self.row],list(app.FIELDS[:-1])+['robot_id'])
    def test_duplicate_ids(self):self.invalid([self.row,self.row])
    def test_unsafe_ids(self):
        for value in ['../R001','CON','r001','R01','R001/x','R0000001']:
            with self.subTest(value=value):self.invalid([[value,*self.row[1:]]])
    def test_invalid_ratings(self):
        for value in ['0','6','NaN','3.5','-1','5e0','']:
            with self.subTest(value=value):self.invalid([[*self.row[:3],value,*self.row[4:]]])
    def test_missing_required_values(self):
        for i in [1,2,6]:
            r=self.row.copy();r[i]=' ';self.invalid([r])
    def test_extra_cells(self):self.invalid([self.row+['unexpected']])
    def test_short_row(self):self.invalid([self.row[:-1]])
    def test_empty_csv(self):self.csv.write_text('');self.assertRaises(ValueError,app.generate_reviews,self.csv,self.out)
    def test_header_only(self):self.invalid([])
    def test_control_characters(self):r=self.row.copy();r[-1]='bad\x01';self.invalid([r])
    def test_long_comment(self):r=self.row.copy();r[-1]='x'*241;self.invalid([r])
    def test_average_rounding(self):self.assertEqual(app.average_rating(dict(zip(app.RATINGS,[5,5,3]))),Decimal('4.33'))
    def test_all_rows_validated_first(self):self.invalid([self.row,['R002',*self.row[1:3],'bad',*self.row[4:]]])
    def test_existing_directory_refused(self):
        self.write();self.out.mkdir();marker=self.out/'keep.txt';marker.write_text('keep')
        self.assertRaises(FileExistsError,app.generate_reviews,self.csv,self.out);self.assertEqual(marker.read_text(),'keep')
    def test_existing_file_refused(self):
        self.write();self.out.write_text('keep');self.assertRaises(FileExistsError,app.generate_reviews,self.csv,self.out)
        self.assertEqual(self.out.read_text(),'keep')
    def test_missing_input(self):self.assertRaises(FileNotFoundError,app.generate_reviews,self.csv,self.out)
    def test_bad_utf8(self):self.csv.write_bytes(b'\xff');self.assertRaises(UnicodeError,app.generate_reviews,self.csv,self.out)
    def test_bom(self):self.write();self.csv.write_bytes(b'\xef\xbb\xbf'+self.csv.read_bytes());self.assertEqual(app.generate_reviews(self.csv,self.out),1)
    def test_cli(self):
        self.write();r=subprocess.run([sys.executable,str(ROOT/'generate_robot_reviews.py'),str(self.csv),'--output',str(self.out)],capture_output=True,text=True)
        self.assertEqual(r.returncode,0,r.stderr);self.assertIn('Reviews generated: 1',r.stdout)
        r=subprocess.run([sys.executable,str(ROOT/'generate_robot_reviews.py'),str(self.csv),'--output',str(self.out)],capture_output=True,text=True)
        self.assertEqual(r.returncode,2);self.assertIn('Error:',r.stderr)
    def test_reproducible_generator(self):
        generate(self.csv);other=self.root/'other.csv';generate(other)
        self.assertEqual(self.csv.read_bytes(),other.read_bytes());self.assertEqual(len(app.load_records(self.csv)),800)
        self.assertRaises(FileExistsError,generate,self.csv)
    def test_800_documents(self):
        generate(self.csv);before=hashlib.sha256(self.csv.read_bytes()).hexdigest()
        self.assertEqual(app.generate_reviews(self.csv,self.out),800)
        files=list(self.out.glob('*.docx'));self.assertEqual(len(files),800)
        records=app.load_records(self.csv)
        for robot in records:
            d=Document(self.out/f"{robot['robot_id']}_review.docx")
            text='\n'.join(p.text for p in d.paragraphs)
            for k in ('robot_id','robot_name','department','manager_comments'):self.assertIn(robot[k],text)
            self.assertEqual(d.tables[0].rows[-1].cells[1].text,str(app.average_rating(robot)))
        self.assertEqual(hashlib.sha256(self.csv.read_bytes()).hexdigest(),before)

if __name__=='__main__':unittest.main()
